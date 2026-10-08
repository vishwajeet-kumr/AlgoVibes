"""
ContractShield AI — Document Parser

Handles PDF and DOCX text extraction.
Uses PyMuPDF (fitz) for PDFs and python-docx for DOCX files.
"""

import io
from dataclasses import dataclass


@dataclass
class ParsedDocument:
    """Result of parsing an uploaded document."""
    raw_text: str
    pages: int
    word_count: int
    filename: str


async def parse_document(file_bytes: bytes, filename: str) -> ParsedDocument:
    """
    Extract text content from a PDF or DOCX file.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename: Original filename (used to determine file type).

    Returns:
        ParsedDocument with extracted text and metadata.

    Raises:
        ValueError: If file type is unsupported or file is corrupted.
    """
    file_ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if file_ext == "pdf":
        return _parse_pdf(file_bytes, filename)
    elif file_ext == "docx":
        return _parse_docx(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported file type: .{file_ext}. Only PDF and DOCX are supported.")


def _parse_pdf(file_bytes: bytes, filename: str) -> ParsedDocument:
    """Extract text from a PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise RuntimeError("PyMuPDF (fitz) is required for PDF parsing. Install with: pip install PyMuPDF")

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Failed to open PDF file: {e}")

    pages_text = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text")
        if text.strip():
            pages_text.append(text)

    doc.close()

    if not pages_text:
        raise ValueError("PDF appears to be empty or contains only images (no extractable text).")

    raw_text = "\n\n".join(pages_text)
    word_count = len(raw_text.split())

    return ParsedDocument(
        raw_text=raw_text,
        pages=len(pages_text),
        word_count=word_count,
        filename=filename,
    )


def _parse_docx(file_bytes: bytes, filename: str) -> ParsedDocument:
    """Extract text from a DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx is required for DOCX parsing. Install with: pip install python-docx")

    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Failed to open DOCX file: {e}")

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    if not paragraphs:
        raise ValueError("DOCX appears to be empty (no extractable text).")

    raw_text = "\n\n".join(paragraphs)
    word_count = len(raw_text.split())

    # DOCX doesn't have a concept of pages like PDF, estimate from content
    # Rough estimate: ~300 words per page
    estimated_pages = max(1, word_count // 300)

    return ParsedDocument(
        raw_text=raw_text,
        pages=estimated_pages,
        word_count=word_count,
        filename=filename,
    )
